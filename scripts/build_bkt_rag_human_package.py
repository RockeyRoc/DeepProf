"""Create blinded, local-only human review files for the M3 BKT/RAG study.

Textbook excerpts and response text are written only to the explicitly selected
private package directory outside the repository. Requires openpyxl for xlsx.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import re
import sqlite3
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT = Path(__file__).resolve().parents[1]
BASE_ID = "m3-abc-bkt-20260928T134542Z-a216df"
SOURCE_AUDIT_REL = Path("bkt-rag-improvement-20260929") / "citation-audit-v2"
QUESTIONS: dict[str, list[tuple[str, str, str]]] = {
    "DS-LIN-02": [
        ("顺序表按下标随机访问是否需要扫描前驱元素？请区分访问与插入成本。", "定义/复杂度", "访问不需要扫描；插入成本取决于后续元素移动"),
        ("长度为 8 的顺序表在第 3 个位置前插入元素，最少需要移动哪些元素？说明计数口径。", "代码/操作", "按 1 起始位置说明需后移的后缀元素数"),
        ("删除顺序表第 i 个元素后，逻辑长度与其后元素位置怎样变化？", "代码/操作", "长度减一，后缀元素前移填补空位"),
        ("已知顺序表长度 n、容量足够，在位置 i 插入；给出最坏移动量并解释 i 的定义。", "复杂度前提", "若 i 为 1 起始位置，移动 n-i+1 个元素；需说明位置定义"),
    ],
    "DS-LIN-04": [
        ("单链表结点 p 后插入新结点 s，应按什么顺序更新 next 指针？", "代码/操作", "先令 s.next=p.next，再令 p.next=s"),
        ("给定带头结点的单链表，要删除 p 的直接后继 q，怎样更新链接并处理 q？", "代码/操作", "p.next=q.next 后释放 q；说明空后继边界"),
        ("在单链表已知尾结点指针时追加一个新结点，哪些链接和尾指针必须更新？", "代码/操作", "旧尾 next 指向新结点，新结点 next 为空，尾指针更新"),
        ("删除目标结点前必须定位其前驱时，查找与改指针的时间成本分别是什么？", "复杂度前提", "定位 O(n)，已知前驱后改链 O(1)；不可把二者混为一谈"),
    ],
    "DS-LIN-05": [
        ("在双向链表结点 p 与 q 之间插入 s，需要维持哪四条前后向链接？", "代码/操作", "s.prev=p,s.next=q,p.next=s,q.prev=s"),
        ("循环链表从头结点遍历时，何时停止才能避免把回到头结点误判为新数据？", "代码/边界", "再到头结点即停止；头结点是否计入数据需区分"),
        ("删除双向链表中的中间结点 q 后，前驱和后继如何重新连接？", "代码/操作", "前驱.next=后继 且 后继.prev=前驱"),
        ("比较单链表与双向链表删除已定位结点时的链接更新及额外存储。", "跨概念", "双向结点多一条 prev，但已定位删除可直接连接两侧"),
    ],
    "DS-LIN-06": [
        ("用栈检查括号串 `([{}])` 是否匹配，逐步说明压入、弹出的条件。", "代码/应用", "左括号入栈，右括号匹配栈顶后弹出，最终栈空"),
        ("将后缀表达式 `23*4+` 按栈求值，列出每一步栈内容。", "代码/操作", "先算 2*3，再加 4，结果 10"),
        ("递归过程为何可用栈描述？说明一次调用至少需要保留哪些返回信息。", "定义/应用", "调用帧后进先出；参数/局部状态/返回位置等"),
        ("一个栈依次执行 push(3), push(5), pop(), push(7), pop()，最后栈顶是什么？", "代码/操作", "剩余 3 为栈顶"),
    ],
    "DS-LIN-07": [
        ("队列依次入队 A、B、C，再出队两次并入队 D，队列从队首到队尾是什么？", "代码/操作", "先出 A、B，余 C 后入 D，队列 C,D"),
        ("循环队列采用留一个空位判满时，front、rear 相等与相邻关系各表示什么？", "定义/边界", "相等为空；(rear+1)%m==front 为满"),
        ("顺序队列 front 不回退时，数组前端已有空位但 rear 到末端；这属于什么现象？", "定义/操作", "假溢出；循环队列可复用前端空间"),
        ("说明用队列实现广度优先遍历时，顶点何时标记访问、何时入队。", "跨概念/代码", "发现时标记并入队，防止重复入队"),
    ],
    "DS-TREE-01": [
        ("在一棵树中，结点度数、叶子结点和森林分别如何定义？", "定义", "子女数；无子结点；互不相交树的集合"),
        ("若结点 u 是 v 的祖先、v 是 w 的父结点，u 与 w 的关系是什么？", "定义/推理", "u 是 w 的祖先；v 是其父结点"),
        ("一棵树有 9 个结点，其中各结点孩子数为 0、1、2 的数量分别是 4、2、3。叶子数是多少？", "计算/定义", "度为 0 的结点共 4 个"),
        ("说明树的结点数与边数关系，并指出适用条件。", "定义/条件", "非空连通无环树 n 个结点有 n-1 条边"),
    ],
    "DS-TREE-03": [
        ("给定二叉树根 A，左子树 B(C,空)，右子树 D(空,E)，写出先序、中序、后序序列。", "代码/遍历", "先根左右；左根右；左右根"),
        ("一棵二叉树的先序为 ABDECFG，中序为 DBEAFCG，如何递归定位根并划分子树？", "跨片段/算法", "先序首项为根；在中序中分割左右子树"),
        ("同一二叉树按层序访问时与先序访问的策略差异是什么？", "对比", "层序按层并用队列；先序用递归/栈先根"),
        ("给定先序 `ABC`、中序 `CBA`，重建后序并说明树形。", "推理/遍历", "根 A；其左子树根 B、左子 C；后序 CBA"),
    ],
    "DS-TREE-04": [
        ("一棵满二叉树有 5 层（根为第 1 层），最多有多少个结点？写出公式。", "复杂度/性质", "2^5-1=31"),
        ("含 7 个叶子的满二叉树有多少个非叶结点？说明前提。", "性质/条件", "每个非叶结点恰有 2 个孩子时 n0=n2+1，因此 6"),
        ("区分完全二叉树与满二叉树，并给出一个只满足前者的形状描述。", "定义/易混", "完全树最后一层靠左可不满；满树所有内部结点两子且叶同层"),
        ("二叉树高度 h 按层数计，最少结点数与最大结点数分别是什么？", "性质/定义", "最少 h（斜链），最多 2^h-1；说明高度口径"),
    ],
    "DS-GRAPH-01": [
        ("含 n 个顶点、m 条边的稀疏图，邻接矩阵与邻接表的空间量级各是什么？", "复杂度前提", "矩阵 O(n^2)；邻接表 O(n+m)"),
        ("把无向边 (u,v) 加入无向图邻接矩阵，需更新几个对称位置？", "代码/表示", "矩阵[u][v] 与 [v][u] 两处"),
        ("有向图邻接表中顶点 v 的出度如何通过表结构得到？", "定义/表示", "统计 v 的出边链表长度"),
        ("稠密图中需要频繁判断任意两点是否相邻，选择哪种表示更直接？说明代价。", "取舍", "邻接矩阵 O(1) 判断，代价是 O(n^2) 空间"),
    ],
    "DS-GRAPH-02": [
        ("无向图边为 A-B,A-C,B-D,C-E,D-F,E-F；从 A 开始，邻接点按字母序，写 BFS 次序。", "代码/遍历", "A,B,C,D,E,F"),
        ("BFS 中某顶点应在何时标记已访问？若出队时才标记会有什么风险？", "代码/边界", "入队时标记；否则多父结点可能重复入队"),
        ("从源点做 BFS 为何可得到无权图最少边数路径？说明按层含义。", "复杂度/正确性", "队列按距离层推进，首次到达即边数最少"),
        ("若图不连通，从一个起点执行 BFS 后哪些顶点未被访问？如何覆盖全图？", "边界/应用", "其他连通分量顶点；对未访问顶点逐一启动"),
    ],
    "DS-GRAPH-03": [
        ("在树形邻接关系 A:{B,C}, B:{D}, C:{E} 中按字母序从 A 做 DFS，写访问序列。", "代码/遍历", "A,B,D,C,E"),
        ("递归 DFS 返回时的访问顺序与进入顶点时记录的顺序有何区别？", "易混/代码", "前者类似后序完成次序；后者为先序发现次序"),
        ("如何用 DFS 判定无向图是否连通？复杂度以邻接表表示说明。", "代码/复杂度", "从一顶点遍历，访问数等于 n；O(n+m)"),
        ("DFS 使用显式栈替代递归时，邻接点压栈顺序如何影响访问顺序？", "代码/前提", "LIFO 导致逆序处理；需反向压栈维持指定次序"),
    ],
    "DS-GRAPH-06": [
        ("用 Dijkstra 求非负权有向图 A→B:2,A→C:6,B→C:1,B→D:5,C→D:1 从 A 的距离。", "代码/算法", "A0,B2,C3,D4"),
        ("Dijkstra 每轮确定当前最小暂定距离顶点后，如何松弛其出边？", "代码/操作", "若 d[u]+w(u,v)<d[v] 则更新 d[v] 与前驱"),
        ("为什么含负权边时标准 Dijkstra 的贪心确定步骤可能错误？", "正确性/条件", "以后续负边可使已确定距离下降；算法假设非负权"),
        ("邻接矩阵实现 Dijkstra 的常见时间复杂度是多少？给出表示与选点前提。", "复杂度前提", "朴素选最小未定点并扫描矩阵时 O(n^2)"),
    ],
    "DS-GRAPH-07": [
        ("给边 A-B:1,A-C:3,B-C:2,B-D:4,C-D:5，写出 Kruskal 选边次序及总权。", "代码/算法", "按权排序并跳过成环边；总权 1+2+4=7"),
        ("Prim 从 A 开始，每轮选择哪类边？它与 Kruskal 的候选边集合差异是什么？", "对比/算法", "跨越已纳入与未纳入顶点的最小边；Kruskal 全局按权选且防环"),
        ("若最小生成树有相同权边，生成树是否一定唯一？说明条件。", "易混/条件", "不一定；唯一性需最小边权选择无歧义等条件"),
        ("一个连通无向图有 n 个顶点，最小生成树必须有多少条边？", "定义/性质", "n-1"),
    ],
    "DS-SORT-01": [
        ("区分稳定排序与原地排序；两者各自讨论的是哪种性质？", "定义/易混", "稳定性保持相等键相对次序；原地关注额外空间"),
        ("说明比较排序 Ω(n log n) 下界依赖的输入/模型假设。", "复杂度前提", "基于比较决策树与需区分 n! 个排列的模型"),
        ("某算法最坏 O(n^2)、平均 O(n log n)。报告其复杂度时还必须说明什么？", "报告条件", "输入分布/枢轴选择/最坏与平均定义"),
        ("稳定、原地、时间复杂度是同一个排序评价维度吗？分别举出报告项。", "方法/定义", "否；顺序保持、额外空间、运行时间分别报告"),
    ],
    "DS-SORT-02": [
        ("对数组 [5,2,4,1] 执行直接插入排序，写出每一趟插入后的数组。", "代码/轨迹", "[2,5,4,1]，[2,4,5,1]，[1,2,4,5]"),
        ("直接插入排序在近乎有序数组上为何可能快于逆序输入？", "复杂度前提", "移动/比较取决于逆序程度，近有序时少移动"),
        ("直接插入排序遇到相等键时，为何插入位置选择会影响稳定性？", "代码/条件", "只在严格大于时右移可保留同键先后"),
        ("对 n 个元素的直接插入排序给出最好、最坏时间量级并说明输入形态。", "复杂度前提", "已有序 Θ(n)，逆序 Θ(n^2)"),
    ],
    "DS-SORT-06": [
        ("将 [8,3,6,2] 分成单元素后执行归并，写出两轮合并结果。", "代码/轨迹", "[3,8]与[2,6]；随后[2,3,6,8]"),
        ("归并排序递归式 T(n)=2T(n/2)+cn 的解是什么？", "复杂度", "Θ(n log n)"),
        ("数组归并排序为何通常需要额外辅助空间？该空间量级是多少？", "复杂度/实现", "合并时暂存元素，常见 O(n) 辅助空间"),
        ("合并两个已排序段时，如何处理相等键才能保持稳定？", "代码/条件", "相等时先取左段元素，保持原相对顺序"),
    ],
    "DS-SORT-07": [
        ("以首元素为枢轴对 [4,1,5,2,3] 做一次划分，给出一种满足分区性质的结果。", "代码/轨迹", "枢轴左侧不大于4，右侧不小于4；需说明具体划分约定"),
        ("输入已升序且每次固定选首元素为枢轴时，快速排序递归规模如何退化？", "复杂度前提", "近似 n-1 与 0，递归深度 Θ(n)，时间 Θ(n^2)"),
        ("随机枢轴快速排序的期望时间量级是什么？答案的期望取决于什么随机性？", "复杂度前提", "期望 Θ(n log n)，随机枢轴选择"),
        ("快速排序一次 partition 完成后，枢轴位置能保证什么？", "正确性/定义", "枢轴最终到位，左右元素分别满足分区关系；左右内部未必有序"),
    ],
}

INSUFFICIENT = [
    ("在动态数组容量翻倍策略下，连续 n 次 append 的总搬移次数是多少？要求给出势能法摊还证明。", "DS-LIN-03", "教材索引明确标为动态数组扩容与摊还复杂度证据缺口；未找到势能法证明。", "dynamic array amortized potential proof"),
    ("在 C11 lock-free 队列中，为何 compare_exchange 需要 acquire-release memory order？", "并发内存模型", "冻结教材不覆盖 C11 原子内存序与无锁队列。", "C11 memory_order acquire release"),
    ("对 B+ 树插入导致叶页分裂的 I/O 次数给出磁盘页模型推导。", "外存 B+ 树", "冻结教材索引未找到 B+ 树页分裂与外存 I/O 模型。", "B+ tree page split I/O model"),
    ("证明红黑树删除后双黑修复的所有旋转情形，并给出摊还高度界。", "红黑树删除", "冻结教材未覆盖红黑树删除修复的完整情形。", "red-black tree deletion double black"),
    ("在 CUDA warp-level primitives 上实现图 BFS，并分析 occupancy 对吞吐的影响。", "GPU 图算法", "冻结教材不包含 CUDA warp、occupancy 或 GPU 并行执行模型。", "CUDA warp occupancy graph BFS"),
    ("给出分布式一致性哈希在节点加入时的数据迁移上界及虚拟节点影响。", "分布式一致性哈希", "冻结教材未覆盖分布式一致性哈希与虚拟节点。", "distributed consistent hashing virtual nodes"),
    ("用 hazard pointers 证明无锁栈的安全内存回收不会发生 use-after-free。", "无锁内存回收", "冻结教材不包含 hazard pointer 内存回收机制。", "hazard pointers memory reclamation"),
    ("估算 Bloom filter 在目标假阳性率 p 下的最优位数组大小和哈希函数数。", "Bloom filter", "冻结教材未找到 Bloom filter 及假阳性率参数化设计。", "Bloom filter optimal bits false positive"),
    ("在 SIMD 指令下并行执行 quicksort partition，分析向量宽度和分支预测的交互。", "SIMD 排序", "冻结教材不包含 SIMD 指令级并行与分支预测分析。", "SIMD quicksort partition branch prediction"),
    ("用 Raft 复制状态机说明图数据库索引更新的一致性保证及线性化点。", "分布式一致性", "冻结教材未覆盖 Raft 与线性化一致性。", "Raft linearizability state machine"),
    ("给出跳表在自适应对抗性插入顺序下的最坏查找界和随机化假设。", "跳表对抗分析", "冻结教材未找到跳表随机化及对抗性最坏界。", "skip list adversarial insertion randomized bound"),
    ("用外存模型证明缓冲树批量更新的摊还 I/O 复杂度。", "缓冲树", "冻结教材未覆盖 buffer tree 与外存摊还 I/O。", "buffer tree amortized I/O"),
    ("为分布式图 partition 设计最小化跨机边的优化目标，并分析通信轮数。", "分布式图计算", "冻结教材不包含分布式图划分与通信轮复杂度。", "distributed graph partition communication rounds"),
    ("解释 CPU cache coherence false sharing 如何改变多线程队列的扩展性。", "缓存一致性", "冻结教材未覆盖多核缓存一致性与 false sharing。", "cache coherence false sharing concurrent queue"),
    ("在密码学抗碰撞假设下设计 Merkle tree 验证路径并证明篡改检测性。", "Merkle tree", "冻结教材未覆盖密码学哈希承诺和 Merkle 证明。", "Merkle tree cryptographic collision resistance"),
    ("对 LSM-tree compaction 推导写放大率与读放大率，并比较 leveled/tiered 策略。", "LSM-tree", "冻结教材未找到 LSM-tree compaction 与放大率分析。", "LSM tree compaction write amplification"),
]

# Keep all covered concepts while meeting the 64-answerable-item split contract.
# The removed items were redundant within their concept family and remain in source control.
QUESTION_TRIMS = {"DS-LIN-05": 3, "DS-TREE-01": 3, "DS-SORT-01": 3, "DS-SORT-07": 3}
DEVELOPMENT_CONCEPTS = {"DS-LIN-02", "DS-LIN-04", "DS-LIN-06", "DS-LIN-07", "DS-TREE-03", "DS-TREE-04", "DS-GRAPH-01", "DS-GRAPH-02"}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def write_csv(path: Path, rows: list[dict[str, Any]], columns: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def blind_id(prefix: str, value: str) -> str:
    return prefix + hashlib.sha256(value.encode("utf-8")).hexdigest()[:12].upper()


def chinese_bigrams(text: str) -> set[str]:
    runs = re.findall(r"[\u4e00-\u9fff]+", text or "")
    values: set[str] = set()
    for run in runs:
        values.update(run[i:i + 2] for i in range(len(run) - 1))
    values.update(token.lower() for token in re.findall(r"[a-zA-Z][a-zA-Z0-9_+-]*|\d+", text or ""))
    return values


def choose_source_excerpt(connection: sqlite3.Connection, concept: dict[str, Any], question: str) -> dict[str, Any]:
    refs = concept.get("source_refs") or []
    module = str(concept.get("module") or "")
    ref_names = [str(ref.get("chapter") or "") for ref in refs]
    candidates = connection.execute(
        "SELECT c.document_id,c.chunk_id,c.page,c.printed_page,c.chapter,c.section,c.text,d.filename,r.title "
        "FROM library_chunks c JOIN library_documents d ON d.document_id=c.document_id "
        "JOIN library_resources r ON r.resource_id=c.resource_id WHERE c.reliable=1"
    ).fetchall()
    query_terms = chinese_bigrams(" ".join([question, str(concept.get("name") or ""),
        str(concept.get("learning_objective") or ""), module, " ".join(ref_names)]))
    module_alias = {"树": "树和二叉树", "排序": "内部排序", "查找": "查找", "线性表": "线性表",
                    "栈和队列": "栈和队列", "图": "图"}.get(module, module)
    ranked: list[tuple[float, sqlite3.Row]] = []
    for row in candidates:
        chapter = str(row[4] or "")
        text = str(row[6] or "")
        section_match = any(name and name in chapter for name in ref_names) or bool(module_alias and module_alias in chapter)
        terms = chinese_bigrams(text)
        overlap = len(query_terms & terms)
        score = overlap + (25 if section_match else 0) + (min(len(text), 1800) / 1800)
        ranked.append((score, row))
    ranked.sort(key=lambda item: (-item[0], int(item[1][2]), str(item[1][1])))
    if not ranked or ranked[0][0] <= 1:
        raise ValueError(f"source_excerpt_not_found_for:{concept.get('concept_id')}")
    row = ranked[0][1]
    text = str(row[6] or "")
    sentence_parts = [part.strip() for part in re.split(r"(?<=[。！？!?；;])\s*|[\r\n]+", text) if part.strip()]
    qterms = chinese_bigrams(question + " " + str(concept.get("name") or ""))
    scored = sorted(((len(qterms & chinese_bigrams(sentence)), i, sentence)
                     for i, sentence in enumerate(sentence_parts)), reverse=True)
    selected = [sentence for score, _, sentence in scored if score >= max(2, (scored[0][0] * 0.5 if scored else 2))][:3]
    excerpt = "\n".join(selected) if selected else text[:900]
    return {"document_id": str(row[0]), "chunk_id": str(row[1]), "pdf_page": int(row[2]),
            "printed_page": row[3], "chapter": str(row[4] or ""), "section": str(row[5] or ""),
            "source_title": str(row[8] or ""), "source_filename": str(row[7] or ""),
            "source_excerpt": excerpt, "evidence_status": "AI-selected source span; human verification required"}


def style_sheet(ws: Any) -> None:
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    if ws.max_row:
        ws.auto_filter.ref = ws.dimensions
    header_fill = PatternFill("solid", fgColor="17365D")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for index, column in enumerate(ws.columns, 1):
        values = [str(cell.value or "") for cell in list(column)[:150]]
        width = min(54, max(12, max((len(value) for value in values), default=12) * 0.85))
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_sheet(workbook: Workbook, name: str, rows: list[dict[str, Any]], columns: list[str]) -> None:
    ws = workbook.create_sheet(title=name)
    ws.append(columns)
    for row in rows:
        ws.append([row.get(column, "") for column in columns])
    style_sheet(ws)


def write_xlsx(path: Path, sheets: list[tuple[str, list[dict[str, Any]], list[str]]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "说明"
    ws.append(["本工作簿用于独立人工评审；不得向评审者提供盲法密钥。"])
    ws.append(["状态：待人工评审。AI 生成/辅助内容均未校准，不作为人工金标。"])
    style_sheet(ws)
    for name, rows, columns in sheets:
        add_sheet(wb, name, rows, columns)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build(source_home: Path, package_root: Path) -> dict[str, Any]:
    base = source_home / "experiments" / "m3-abc-research" / BASE_ID
    audit_root = base / SOURCE_AUDIT_REL
    private_root = package_root.resolve()
    private_root.mkdir(parents=True, exist_ok=True)
    citation_source = audit_root / "review-worksheet-ai-reviewed.csv"
    retrieval_source = audit_root / "rag-relevance-worksheet.csv"
    course_retrieval_source = source_home / "experiments" / "m3-abc-research" / "bkt-rag-improvement-20261001" / "course-retrieval" / "course-rag-method-candidates.csv"
    citation_rows = read_csv(citation_source)
    retrieval_rows = read_csv(retrieval_source)
    historical_unique_pairs = len({(row.get("case_id", ""), row.get("chunk_id", "")) for row in retrieval_rows})
    course_retrieval_rows = read_csv(course_retrieval_source) if course_retrieval_source.exists() else []

    citation_key: list[dict[str, Any]] = []
    blind_claim: list[dict[str, Any]] = []
    for row in citation_rows:
        raw_key = "|".join([row.get("cell_id", ""), row.get("claim_id", ""),
            row.get("citation_selector", ""), str(row.get("locator", ""))])
        item_id = blind_id("C-", raw_key)
        citation_key.append({"item_id": item_id, "cell_id": row.get("cell_id", ""),
            "claim_id": row.get("claim_id", ""), "citation_selector": row.get("citation_selector", ""),
            "locator": row.get("locator", ""), "record_type": row.get("record_type", "")})
        blind_claim.append({"item_id": item_id, "blind_case_id": blind_id("Q-", row.get("case_id", "")),
            "claim_text": row.get("claim_text", ""), "answer_citation_selector": row.get("citation_selector", ""),
            "document_title": row.get("source_title", ""), "source_file": row.get("source_file", ""),
            "pdf_page": row.get("source_page", ""), "printed_page": row.get("printed_page", ""),
            "chapter": row.get("chapter", ""), "section": row.get("section", ""),
            "source_excerpt": row.get("source_excerpt", ""), "locator_status_for_review": "",
            "individual_support_label": "", "individual_reason": "", "joint_support_label": "",
            "joint_reason": "", "reviewer": ""})
    cited_pairs = [row for row in blind_claim if next((src.get("record_type") for src in citation_key
                  if src["item_id"] == row["item_id"]), "") == "claim_citation_pair"]
    pilot_citation_ids = {row["item_id"] for row in sorted(cited_pairs,
        key=lambda item: hashlib.sha256(("citation-pilot:" + item["item_id"]).encode()).hexdigest())[:20]}
    for row in blind_claim:
        row["pilot_sample"] = "yes" if row["item_id"] in pilot_citation_ids else "no"
    citation_columns = ["item_id", "blind_case_id", "claim_text", "answer_citation_selector",
        "document_title", "source_file", "pdf_page", "printed_page", "chapter", "section",
        "source_excerpt", "locator_status_for_review", "individual_support_label", "individual_reason",
        "joint_support_label", "joint_reason", "reviewer", "pilot_sample"]
    for rater in (1, 2):
        ordered = sorted(blind_claim, key=lambda item: hashlib.sha256(
            f"citation-rater-{rater}:{item['item_id']}".encode()).hexdigest())
        write_csv(private_root / f"citation-rater-{rater}.csv", ordered, citation_columns)
    citation_adjudication = [{"item_id": row["item_id"], "claim_text": row["claim_text"],
        "source_excerpt": row["source_excerpt"], "rater_1_label": "", "rater_1_reason": "",
        "rater_2_label": "", "rater_2_reason": "", "adjudicated_label": "",
        "adjudication_reason": "", "reviewer": ""} for row in blind_claim]
    citation_adjudication_columns = list(citation_adjudication[0]) if citation_adjudication else []
    write_csv(private_root / "citation-adjudication.csv", citation_adjudication, citation_adjudication_columns)

    # Keep the historical candidate pool and append only novel case/chunk pairs
    # from the offline BGE/M3 runs. Method ranks and scores stay in the private key.
    merged_candidates: dict[tuple[str, str], dict[str, Any]] = {}
    for row in retrieval_rows:
        key = (row.get("case_id", ""), row.get("chunk_id", ""))
        if key not in merged_candidates:
            merged_candidates[key] = {**row, "_course_methods": []}
    for row in course_retrieval_rows:
        key = (row.get("case_id", ""), row.get("chunk_id", ""))
        if not all(key):
            continue
        if key not in merged_candidates:
            merged_candidates[key] = {
                "case_id": row.get("case_id", ""), "chunk_id": row.get("chunk_id", ""),
                "expected_no_evidence": row.get("expected_no_evidence", ""),
                "query": row.get("query", ""), "source": row.get("source", ""),
                "document_id": row.get("document_id", ""), "page": row.get("page", ""),
                "printed_page": row.get("printed_page", ""), "chapter": row.get("chapter", ""),
                "section": row.get("section", ""), "source_excerpt": row.get("excerpt", ""),
                "hash_dense_rank": "", "bm25_rank": "", "rrf_rank": "",
                "dense_score": "", "bm25_score": "", "rrf_score": "",
                "total_retrieval_latency_ms": "", "_course_methods": []}
        merged_candidates[key]["_course_methods"].append({
            "method": row.get("method", ""), "rank": row.get("rank", ""),
            "score": row.get("score", ""), "index_config_id": row.get("index_config_id", ""),
            "model_revision": row.get("model_revision", "")})
    retrieval_rows = list(merged_candidates.values())
    old_retrieval_pair_count = historical_unique_pairs
    course_retrieval_added_pairs = sum(1 for row in retrieval_rows if row.get("_course_methods") and
        (row.get("hash_dense_rank") == "" and row.get("bm25_rank") == "" and row.get("rrf_rank") == ""))

    retrieval_key: list[dict[str, Any]] = []
    blind_retrieval: list[dict[str, Any]] = []
    for row in retrieval_rows:
        raw_key = "|".join([row.get("case_id", ""), row.get("chunk_id", "")])
        item_id = blind_id("R-", raw_key)
        retrieval_key.append({"item_id": item_id, "case_id": row.get("case_id", ""),
            "chunk_id": row.get("chunk_id", ""), "expected_no_evidence": row.get("expected_no_evidence", ""),
            "hash_dense_rank": row.get("hash_dense_rank", ""), "bm25_rank": row.get("bm25_rank", ""),
            "rrf_rank": row.get("rrf_rank", ""), "dense_score": row.get("dense_score", ""),
            "bm25_score": row.get("bm25_score", ""), "rrf_score": row.get("rrf_score", ""),
            "latency": row.get("total_retrieval_latency_ms", ""),
            "course_method_ranks_scores": json.dumps(row.get("_course_methods", []), ensure_ascii=False)})
        blind_retrieval.append({"item_id": item_id, "blind_query_id": blind_id("Q-", row.get("case_id", "")),
            "query_text": row.get("query", ""), "document_title": row.get("source", ""),
            "document_id": row.get("document_id", ""), "chunk_id": row.get("chunk_id", ""),
            "pdf_page": row.get("page", ""), "printed_page": row.get("printed_page", ""),
            "chapter": row.get("chapter", ""), "section": row.get("section", ""),
            "candidate_excerpt": row.get("source_excerpt", row.get("excerpt", "")), "relevance_grade_0_1_2": "",
            "rationale": "", "reviewer": ""})
    per_case: dict[str, list[dict[str, Any]]] = {}
    for row in blind_retrieval:
        per_case.setdefault(row["blind_query_id"], []).append(row)
    pilot_retrieval_ids: set[str] = set()
    pilot_queries = sorted(per_case, key=lambda value: hashlib.sha256(("retrieval-pilot:" + value).encode()).hexdigest())[:10]
    for query_id in pilot_queries:
        candidates = sorted(per_case[query_id], key=lambda item: hashlib.sha256(
            ("retrieval-candidate:" + item["item_id"]).encode()).hexdigest())[:3]
        pilot_retrieval_ids.update(item["item_id"] for item in candidates)
    for row in blind_retrieval:
        row["pilot_sample"] = "yes" if row["item_id"] in pilot_retrieval_ids else "no"
    retrieval_columns = ["item_id", "blind_query_id", "query_text", "document_title", "document_id",
        "chunk_id", "pdf_page", "printed_page", "chapter", "section", "candidate_excerpt",
        "relevance_grade_0_1_2", "rationale", "reviewer", "pilot_sample"]
    for rater in (1, 2):
        ordered = sorted(blind_retrieval, key=lambda item: hashlib.sha256(
            f"retrieval-rater-{rater}:{item['item_id']}".encode()).hexdigest())
        write_csv(private_root / f"retrieval-rater-{rater}.csv", ordered, retrieval_columns)
    retrieval_adjudication = [{"item_id": row["item_id"], "query_text": row["query_text"],
        "candidate_excerpt": row["candidate_excerpt"], "rater_1_grade": "", "rater_1_reason": "",
        "rater_2_grade": "", "rater_2_reason": "", "adjudicated_grade": "",
        "adjudication_reason": "", "reviewer": ""} for row in blind_retrieval]
    retrieval_adjudication_columns = list(retrieval_adjudication[0]) if retrieval_adjudication else []
    write_csv(private_root / "retrieval-adjudication.csv", retrieval_adjudication, retrieval_adjudication_columns)

    manifest_path = PROJECT / "data" / "courses" / "data_structures_c" / "manifest.json"
    course = json.loads(manifest_path.read_text(encoding="utf-8"))
    concepts = {row["concept_id"]: row for row in course.get("concepts", [])}
    database = source_home / "experiments" / "m3-abc-research" / BASE_ID / "runtime" / "sessions.sqlite"
    db = sqlite3.connect(database.resolve().as_uri() + "?mode=ro", uri=True)
    db.row_factory = sqlite3.Row
    answerable: list[dict[str, Any]] = []
    serial = 1
    answerable_ids = list(QUESTIONS)
    for concept_id in answerable_ids:
        concept = concepts[concept_id]
        kept = [(index, item) for index, item in enumerate(QUESTIONS[concept_id])
                if index != QUESTION_TRIMS.get(concept_id, -1)]
        for index, (question, category, conditions) in kept:
            split = "development" if concept_id in DEVELOPMENT_CONCEPTS else "sealed_test"
            evidence = choose_source_excerpt(db, concept, question)
            answerable.append({"item_id": f"NQ-{serial:03d}", "question_family_id": f"F-{concept_id}",
                "split": split, "concept_id": concept_id, "concept_name": concept.get("name", ""),
                "question_category": category, "question_text": question, "required_conditions": conditions,
                "draft_answerability": "answerable", "gold_status": "draft_pending_human_confirmation",
                "standard_evidence_text": evidence["source_excerpt"],
                "document_id": evidence["document_id"], "chunk_id": evidence["chunk_id"],
                "pdf_page": evidence["pdf_page"], "printed_page": evidence["printed_page"],
                "chapter": evidence["chapter"], "section": evidence["section"],
                "source_title": evidence["source_title"], "evidence_status": evidence["evidence_status"],
                "source_ref_ranges": json.dumps(concept.get("source_refs") or [], ensure_ascii=False)})
            serial += 1
    insufficient: list[dict[str, Any]] = []
    for index, (question, concept, gap, search_terms) in enumerate(INSUFFICIENT):
        split = "development" if index < 8 else "sealed_test"
        insufficient.append({"item_id": f"NQ-{serial:03d}", "question_family_id": f"F-OUT-{index + 1:02d}",
            "split": split, "concept_id": concept, "concept_name": concept,
            "question_category": "evidence_insufficient", "question_text": question,
            "required_conditions": "按冻结教材原文给出结论；若无直接证据请明确指出缺口。",
            "draft_answerability": "insufficient_evidence", "gold_status": "draft_pending_human_confirmation",
            "standard_evidence_text": "", "document_id": "", "chunk_id": "", "pdf_page": "",
            "printed_page": "", "chapter": "", "section": "", "source_title": "",
            "evidence_status": "no direct span preselected; human must verify against the complete frozen book",
            "evidence_gap_reason": gap, "search_terms": search_terms})
        serial += 1
    db.close()
    questions = answerable + insufficient
    if len(questions) != 80:
        raise ValueError(f"expected_80_questions_got_{len(questions)}")
    counts: dict[str, dict[str, int]] = {}
    for split in ("development", "sealed_test"):
        rows = [row for row in questions if row["split"] == split]
        counts[split] = {"total": len(rows), "answerable": sum(row["draft_answerability"] == "answerable" for row in rows),
                         "insufficient": sum(row["draft_answerability"] == "insufficient_evidence" for row in rows)}
        if counts[split] != {"total": 40, "answerable": 32, "insufficient": 8}:
            raise ValueError(f"question_split_count_mismatch:{split}:{counts[split]}")
    families: dict[str, set[str]] = {}
    for row in questions:
        families.setdefault(row["question_family_id"], set()).add(row["split"])
    if any(len(value) != 1 for value in families.values()):
        raise ValueError("question_family_crosses_splits")
    public_question_columns = list(questions[0])
    write_csv(private_root / "new-80-question-owner.csv", questions, public_question_columns)
    question_key = [{"item_id": row["item_id"], "split": row["split"],
        "draft_answerability": row["draft_answerability"], "concept_id": row["concept_id"],
        "question_family_id": row["question_family_id"], "gold_status": row["gold_status"]} for row in questions]
    write_csv(private_root / "new-80-sealed-key.csv", question_key, list(question_key[0]))
    blind_questions = [{"item_id": blind_id("N-", row["item_id"]), "question_text": row["question_text"],
        "required_conditions": row["required_conditions"], "human_answerability": "",
        "human_evidence_location": "", "rationale": "", "reviewer": ""} for row in questions]
    question_blind_cols = list(blind_questions[0])
    for rater in (1, 2):
        ordered = sorted(blind_questions, key=lambda row: hashlib.sha256(
            f"question-rater-{rater}:{row['item_id']}".encode()).hexdigest())
        write_csv(private_root / f"new-80-rater-{rater}.csv", ordered, question_blind_cols)
    question_adjudication = [{"item_id": row["item_id"], "question_text": row["question_text"],
        "rater_1_answerability": "", "rater_1_evidence_location": "", "rater_1_reason": "",
        "rater_2_answerability": "", "rater_2_evidence_location": "", "rater_2_reason": "",
        "adjudicated_answerability": "", "adjudication_reason": "", "reviewer": ""} for row in blind_questions]
    write_csv(private_root / "new-80-adjudication.csv", question_adjudication, list(question_adjudication[0]))

    write_csv(private_root / "citation-blinding-key.csv", citation_key, list(citation_key[0]))
    write_csv(private_root / "retrieval-blinding-key.csv", retrieval_key, list(retrieval_key[0]))
    source_hashes = {"annotation_builder_script": sha256(Path(__file__).resolve()), "citation_ai_review_csv": sha256(citation_source), "retrieval_candidate_csv": sha256(retrieval_source),
                     "frozen_textbook_index": json.loads((PROJECT / "docs" / "experiments" / "bkt-rag-improvement-20260929" / "rag-relevance-worksheet-manifest.json").read_text(encoding="utf-8"))["frozen_library_database_sha256"],
                     "course_manifest": sha256(manifest_path), "frozen_sqlite_path": str(database),
                     "frozen_sqlite_sha256": hashlib.sha256(database.read_bytes()).hexdigest()}
    if course_retrieval_source.exists():
        source_hashes["course_rag_method_candidates_csv"] = sha256(course_retrieval_source)
    codebook = [
        {"task": "引用—主张配对", "label": label, "definition": definition}
        for label, definition in [("完整支持", "该来源单独覆盖主张的全部关键事实与条件。"),
            ("部分支持", "来源支持部分内容，但缺少至少一个关键事实或前提。"),
            ("无支持", "无相关证据、仅主题相似或关键词重合。"),
            ("矛盾", "来源明确否定主张中的至少一项事实。"),
            ("无法判定", "位置/文本/问题有歧义，或证据文本无法可靠判读。")]]
    codebook += [{"task": "检索候选相关性", "label": str(i), "definition": text} for i, text in [
        (0, "不相关或只有主题相似"), (1, "部分支持，但缺少关键证据"),
        (2, "直接支持问题的关键证据需求")]]
    write_csv(private_root / "codebook.csv", codebook, ["task", "label", "definition"])
    guide = """# M3 人工标注包使用指南\n\n本目录只在项目授权的本地评审环境保存回答、教材片段和题库草案；不得把教材摘录合并进公开仓库或重新分发。盲评者只收对应的 `citation-rater-N.csv`、`retrieval-rater-N.csv`、`new-80-rater-N.csv` 及对应工作簿，不收 `*-key.csv`、owner 表或另一位评审者的文件。\n\n## 评审规则\n\n1. 先各自独立完成 20 个引用配对与 30 个检索候选试标；本包以 `pilot_sample=yes` 标记。讨论规则后冻结 codebook，再评剩余项目。\n2. 引用按每个“事实主张—回答实际引用”配对判断。字段能解析、候选曾被检索到、决策曾附带引用，都不能代替语义支持。只有覆盖主张的全部关键事实和前提才标“完整支持”。\n3. 多引用的单条配对分别标注；再独立判断联合证据是否完整。不得把两个“部分支持”自动加成“完整支持”。\n4. 没有回答可见引用的事实主张，主张证据覆盖标“无支持”；不把它计入主张—引用配对分母。找不到来源/页码时标“无法判定”并说明。\n5. 检索标签 0/1/2 分别代表不相关、部分支持、直接满足关键证据需求。只看问题和候选原文，不推断未展示候选池之外的召回。\n6. 新 80 题是 AI 起草包。两位评审要在完整冻结教材中独立核验可答性、必要前提和标准证据位置。AI 提议的划分与证据跨度不是金标。封存集标签在裁决前由数据管理员保管。\n7. 试标后讨论争议例，修订规则版本；正式评审互相隔离。完成后将 CSV 交回给管理员，通过 `scripts/validate_bkt_rag_annotations.py` 检查必填项、重复项、一致率和 κ。\n\n## 报告一致性\n\n引用标签使用原始一致率及五分类 Cohen’s κ。检索 0/1/2 标签同时报告原始一致率及二次加权 Cohen’s κ。报告分母、缺失数和评审者；不以 κ 替代分歧裁决。分歧由第三位裁决者查看双方理由、原始回答和原文位置后给出最终标签。\n\nARES 的上下文相关性、忠实性和答案相关性维度仅作设计参考；本包没有 ARES 人工校准，也不构成其评分。\n"""
    (private_root / "标注指南.md").write_text(guide, encoding="utf-8")

    # Three separate reviewer workbooks preserve blinding: no keys or AI labels are included.
    for rater in (1, 2):
        write_xlsx(private_root / f"human-review-rater-{rater}.xlsx", [
            ("引用", read_csv(private_root / f"citation-rater-{rater}.csv"), citation_columns),
            ("检索", read_csv(private_root / f"retrieval-rater-{rater}.csv"), retrieval_columns),
            ("新题", read_csv(private_root / f"new-80-rater-{rater}.csv"), question_blind_cols),
            ("codebook", codebook, ["task", "label", "definition"]),
        ])
    write_xlsx(private_root / "human-adjudication.xlsx", [
        ("引用裁决", citation_adjudication, citation_adjudication_columns),
        ("检索裁决", retrieval_adjudication, retrieval_adjudication_columns),
        ("新题裁决", question_adjudication, list(question_adjudication[0])),
        ("codebook", codebook, ["task", "label", "definition"]),
    ])

    manifest = {"schema_version": "deepprof-m3-human-annotation-package-v1",
        "status": "ready_for_human_review; all new question gold labels pending",
        "private_local_only": True, "private_package_root": str(private_root),
        "counts": {"citation_rows": len(blind_claim), "citation_pair_rows": len(cited_pairs),
            "citation_pilot_rows": len(pilot_citation_ids), "retrieval_rows": len(blind_retrieval),
            "retrieval_historical_unique_pairs": old_retrieval_pair_count,
            "retrieval_course_method_rows": len(course_retrieval_rows),
            "retrieval_course_new_unique_pairs": course_retrieval_added_pairs,
            "retrieval_pilot_rows": len(pilot_retrieval_ids), "new_questions": len(questions),
            "new_question_splits": counts, "new_question_family_count": len(families)},
        "sources": source_hashes, "files": {},
        "limitations": ["Two existing blind reviews remain pending and are not merged.",
            "Question and evidence spans are AI-generated candidates pending independent confirmation.",
            "Source excerpts are private and not copied into repository output.",
            "No inter-rater statistic exists until completed files are imported."]}
    for path in sorted(private_root.iterdir()):
        if path.is_file() and path.name != "package-manifest.json":
            manifest["files"][path.name] = {"sha256": sha256(path), "bytes": path.stat().st_size}
    (private_root / "package-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-home", type=Path, default=PROJECT.parent / "开发者测试.deepprof")
    parser.add_argument("--package-root", type=Path, required=True,
        help="Private local package path outside the DeepProf repository")
    args = parser.parse_args()
    result = build(args.source_home, args.package_root)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
