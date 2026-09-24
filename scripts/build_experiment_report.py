"""Build the Markdown chart library and editable DOCX from one snapshot."""
from __future__ import annotations

import argparse
import csv
import json
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs" / "experiments"
SOURCE = DOCS / "source-data"
SNAPSHOT_PATH = SOURCE / "experiment-snapshot.json"
CSV_PATH = SOURCE / "figure-data.csv"
MD_PATH = DOCS / "\u56fe\u8868\u5e93.md"
DOCX_PATH = DOCS / "\u56fe\u8868\u5e93.docx"
ANNOTATION_PATH = DOCS / "\u56fe\u8868\u4eba\u5de5\u6ce8\u91ca.md"
FIGURES = {
    "F01": "F01-question-bank",
    "F02": "F02-ab-acceptance",
    "F03": "F03-textbook-page-audit",
    "F04": "F04-constructed-case-suite",
    "F05": "F05-ab-runtime-overview",
}
CASE_NAMES = {
    "cold_start": "\u51b7\u542f\u52a8",
    "prior_insufficient": "\u5148\u9a8c\u4e0d\u8db3",
    "consecutive_errors": "\u8fde\u7eed\u9519\u8bef",
    "hint_then_success": "\u63d0\u793a\u540e\u6210\u529f",
    "insufficient_evidence": "\u8bc1\u636e\u4e0d\u8db3",
    "misconception": "\u6982\u5ff5\u8bef\u89e3",
    "active_explanation_request": "\u4e3b\u52a8\u8bb2\u89e3",
    "low_progress": "\u591a\u8f6e\u65e0\u8fdb\u5c55",
}


def load_snapshot() -> dict[str, Any]:
    return json.loads(SNAPSHOT_PATH.read_text(encoding="utf-8"))


def read_csv() -> dict[str, list[dict[str, Any]]]:
    result: dict[str, list[dict[str, Any]]] = defaultdict(list)
    with CSV_PATH.open(encoding="utf-8", newline="") as stream:
        for row in csv.DictReader(stream):
            row["count"] = int(row["count"])
            row["denominator"] = int(row["denominator"])
            result[row["figure_id"]].append(row)
    return result


def verify(snapshot: dict[str, Any], rows: dict[str, list[dict[str, Any]]]) -> None:
    expected = {
        "F01": snapshot["question_bank"]["question_count"],
        "F02": snapshot["real_ab_acceptance"]["planned_cells"],
        "F03": snapshot["textbook_page_audit"]["total_pdf_pages"],
        "F04": snapshot["constructed_cases"]["total"],
        "F05": snapshot["real_ab_acceptance"]["planned_cells"],
    }
    for figure_id, total in expected.items():
        observed = sum(row["count"] for row in rows[figure_id])
        if figure_id == "F05":
            observed = sum(row["count"] for row in rows[figure_id] if row["status"] == "completed")
            denominators = {row["category"]: row["denominator"] for row in rows[figure_id]}
            if observed != snapshot["real_ab_acceptance"]["completed_cells"] or sum(denominators.values()) != total:
                raise ValueError("F05 completion counts or denominators do not match the experiment snapshot.")
            continue
        if observed != total:
            raise ValueError(f"{figure_id} chart total does not match the experiment snapshot.")
    if not ANNOTATION_PATH.is_file():
        raise FileNotFoundError("The append-only manual annotation file is missing.")
    for figure_id, stem in FIGURES.items():
        for suffix in ("png", "svg", "pdf", "tiff"):
            if not (DOCS / "figures" / f"{stem}.{suffix}").is_file():
                raise FileNotFoundError(f"Missing {figure_id} {suffix} export.")
    if not (SOURCE / "turn-performance.csv").is_file():
        raise FileNotFoundError("The redacted per-turn performance source is missing.")


def plot_versions() -> str:
    lines = SOURCE.joinpath("plot-runtime.txt").read_text(encoding="utf-8").splitlines()
    return ", ".join(line for line in lines if line.startswith(("R ", "ggplot2 ", "patchwork ", "svglite ", "ragg ")))


def fmt(value: Any) -> str:
    return "\u2014" if value is None else str(value)


def acceptance_qa_status_line(snapshot: dict[str, Any]) -> str:
    audit = snapshot.get("acceptance_audit") or {}
    validation = snapshot.get("validation") or {}
    fault = validation.get("fault_injection") or {}
    if not isinstance(fault, dict):
        fault = {}
    recovery = validation.get("process_recovery", "not_run")
    return (
        f"\u81ea\u52a8\u6848\u4f8b\u884c\u4e3a\u6838\u9a8c {audit.get('status', '\u672a\u8fd0\u884c')}\uff1b"
        f"\u6559\u5e08\u8bc4\u5206\u3001\u8bed\u4e49\u6cc4\u9732\u548c\u5f15\u7528\u652f\u6301\u5ea6\u5f85\u4eba\u5de5\u8bc4\u4f30\uff1b"
        f"\u6545\u969c\u6ce8\u5165 {fault.get('status', 'not_run')}\uff08{fault.get('tests_passed', 0)} \u9879\uff09\uff1b"
        f"\u8fdb\u7a0b\u6062\u590d {recovery}\u3002"
    )


def automation_validation_line(snapshot: dict[str, Any]) -> str:
    validation = snapshot.get("validation") or {}
    python = validation.get("python") or {}
    cli = validation.get("cli") or {}
    return (
        f"自动化回归：Python {python.get('passed', '未记录')} 项通过 / "
        f"{python.get('skipped', '—')} 项跳过；CLI {cli.get('tests_passed', '未记录')} 项通过；"
        f"类型检查 {cli.get('typecheck', '未记录')}。"
    )


def descriptions(snapshot: dict[str, Any]) -> dict[str, tuple[str, str, str]]:
    ab = snapshot["real_ab_acceptance"]
    page = snapshot.get("textbook_page_audit", {
        "manually_verified_pages": 0, "pending_manual_pages": 347, "total_pdf_pages": 347,
    })
    run_id = ab["run_id"] or "\u65e0"
    audit = snapshot.get("acceptance_audit") or {}
    audit_status = audit.get("status", "\u672a\u6267\u884c")
    group_text = ", ".join(
        f"{group} {snapshot.get('performance_by_group', {}).get(group, {}).get('latency_ms', {}).get('n', 0)} \u56de\u5408"
        for group in ("A", "B")
    )
    token_group_text = ", ".join(
        f"{group} {snapshot.get('performance_by_group', {}).get(group, {}).get('token_usage', {}).get('n', 0)} \u56de\u5408"
        for group in ("A", "B")
    )
    return {
        "F01": (
            "\u9898\u5e93\u56db\u4e2a\u6a21\u5757\u6309\u53ef\u7528\u8bc4\u5206\u72b6\u6001\u5806\u53e0\u663e\u793a\u3002",
            "\u5206\u5b50/\u5206\u6bcd\u4e3a\u5404\u72b6\u6001\u9898\u6570 / \u672c\u6a21\u5757\u9898\u6570\uff1b\u5355\u4f4d\u9053\uff1bN=30\u3002",
            "\u9898\u5e93 30 \u9898\uff1b16 \u9898\u53ef\u4f7f\u7528\u786e\u5b9a\u6027\u8bc4\u5206\uff0c14 \u9898\u4eba\u5de5\u8bc4\u5206\u6216\u672a\u542f\u7528\u3002",
        ),
        "F02": (
            "40 \u4e2a\u51bb\u7ed3\u6784\u9020\u6848\u4f8b\u8ba1\u5212\u5728 A \u548c B \u4e24\u7ec4\u5404\u8fd0\u884c\u4e00\u6b21\u3002",
            "\u5206\u5b50/\u5206\u6bcd\u4e3a\u6bcf\u7ec4\u5b8c\u6210\u683c\u6570 / 40 \u4e2a\u8ba1\u5212\u683c\uff1b\u5355\u4f4d\u6848\u4f8b\u8fd0\u884c\u683c\uff1b\u8ba1\u5212 N=80\u3002",
            f"run_id={run_id}\uff1b\u5b8c\u6210 {ab['completed_cells']}/{ab['planned_cells']} \u683c\uff1b\u81ea\u52a8\u6848\u4f8b\u6838\u9a8c {audit_status}\uff1b\u672a\u4f7f\u7528 Fake Provider\u3002",
        ),
        "F03": (
            "\u914d\u5957\u6559\u6750\u9875\u7801\u6838\u9a8c\u8fdb\u5ea6\u3002",
            f"\u5206\u5b50/\u5206\u6bcd\u4e3a\u5df2\u6838\u9875\u6570\u3001\u5f85\u6838\u9875\u6570 / {page['total_pdf_pages']} \u9875 PDF\uff1b\u5355\u4f4d PDF \u9875\uff1bN={page['total_pdf_pages']}\u3002",
            f"\u5f53\u524d {page['manually_verified_pages']} \u9875\u5df2\u4eba\u5de5\u786e\u8ba4\uff0c{page['pending_manual_pages']} \u9875\u5f85\u6838\u3002\u672a\u4f7f\u7528\u7edf\u4e00\u9875\u7801\u504f\u79fb\u3002",
        ),
        "F04": (
            "\u51bb\u7ed3\u7684\u516b\u7c7b\u5f00\u53d1\u6848\u4f8b\u6bcf\u7c7b\u5404 5 \u4f8b\u3002",
            "\u5206\u5b50/\u5206\u6bcd\u4e3a\u6bcf\u7c7b\u6848\u4f8b\u6570 / 40 \u4e2a\u6784\u9020\u6848\u4f8b\uff1b\u5355\u4f4d\u4f8b\uff1bN=40\u3002",
            "\u4ec5\u7528\u4e8e\u7a0b\u5e8f\u5206\u652f\u9a8c\u8bc1\uff0c\u4e0d\u662f\u5b66\u751f\u6216\u771f\u4eba\u53c2\u4e0e\u8005\u6570\u636e\u3002",
        ),
        "F05": (
            "A/B \u4e24\u7ec4\u6848\u4f8b\u5b8c\u6210\u60c5\u51b5\u3001\u811a\u672c\u5316\u56de\u5408\u5ef6\u8fdf\u4e0e Token \u5206\u5e03\u3002",
            f"\u5b8c\u6210\u683c N={ab['completed_cells']}/{ab['planned_cells']}\uff1b\u5ef6\u8fdf\u6837\u672c {group_text}\uff1bToken \u6837\u672c {token_group_text}\uff1b\u5355\u4f4d\u5206\u522b\u4e3a\u6848\u4f8b\u683c\u3001ms/\u56de\u5408\u548c Token/\u56de\u5408\u3002",
            f"\u56fe\u4e2d\u70b9\u8868\u793a\u6bcf\u4e2a\u811a\u672c\u5316\u6a21\u578b\u56de\u5408\uff1brun_id={run_id}\u3002\u8fd9\u662f\u6784\u9020\u6848\u4f8b\u7684\u5de5\u7a0b\u8fd0\u884c\u6d4b\u91cf\uff0c\u4e0d\u8868\u793a\u5b66\u4e60\u6210\u6548\u3002",
        ),
    }


def make_markdown(snapshot: dict[str, Any], data: dict[str, list[dict[str, Any]]]) -> str:
    bank = snapshot["question_bank"]
    ab = snapshot["real_ab_acceptance"]
    perf = snapshot["performance"]
    page = snapshot["textbook_page_audit"]
    cases = snapshot["constructed_cases"]
    code = snapshot["code"]
    validation = snapshot.get("validation", {})
    audit = snapshot.get("acceptance_audit") or {}
    coverage = Counter(row.get("coverage_status", "unknown") for row in snapshot.get("course_coverage", {}).get("coverage", []))
    uncovered_concepts = [row.get("concept_id", "") for row in snapshot.get("course_coverage", {}).get("coverage", [])
                          if row.get("coverage_status") in {"uncovered", "not_covered"}]
    runtime = plot_versions()
    lines = [
        "# DeepProf \u9a8c\u6536\u4e0e\u5b9e\u9a8c\u56fe\u8868\u5e93", "",
        f"> \u5feb\u7167 {snapshot['snapshot_version']} \u00b7 {snapshot['generated_at']} \u00b7 \u56fe\u8868\u73af\u5883 {runtime}", "",
        "\u672c\u6587\u6863\u4e0e [\u56fe\u8868\u5e93.docx](\u56fe\u8868\u5e93.docx) \u4f7f\u7528\u540c\u4e00\u4efd\u5feb\u7167\u3002\u4eba\u5de5\u8bf4\u660e\u4fdd\u5b58\u5728 [\u56fe\u8868\u4eba\u5de5\u6ce8\u91ca.md](\u56fe\u8868\u4eba\u5de5\u6ce8\u91ca.md)\uff0c\u91cd\u5efa\u62a5\u544a\u4e0d\u4f1a\u8986\u76d6\u6ce8\u91ca\u3002", "",
        "## \u9a8c\u6536\u72b6\u6001", "",
        "| \u9879\u76ee | \u7ed3\u679c | \u8bf4\u660e |", "|---|---:|---|",
        f"| M0 \u9898\u5e93 | {bank['question_count']} \u9898 | {bank['verified_source_count']} \u9898\u6765\u6e90\u5df2\u76ee\u89c6\u6838\u5bf9\uff1b{bank['pending_reference_count']} \u9898\u5f15\u7528\u5f85\u6838 |",
        f"| \u786e\u5b9a\u6027\u8bc4\u5206 | {bank['deterministic_count']} / {bank['question_count']} | \u4ec5\u542f\u7528\u5df2\u590d\u6838\u89c4\u5219 |",
        f"| \u6559\u5e08\u5ba1\u6838 | {bank['teacher_approval']} | \u5f85\u6559\u5e08\u72ec\u7acb\u5ba1\u6838 |",
        f"| M1 A/B | {ab['completed_cells']} / {ab['planned_cells']} \u683c | {ab['status']}\uff1bProvider {ab.get('provider_profile') or '\u672a\u914d\u7f6e'} / {ab.get('model') or '\u672a\u914d\u7f6e'}\uff1b\u884c\u4e3a\u6838\u9a8c {audit.get('status', '\u672a\u6267\u884c')} |",
        f"| M1 \u81ea\u52a8\u884c\u4e3a | {audit.get('automated_checks', {}).get('cells_passed', 0)} / {ab['planned_cells']} \u683c\u901a\u8fc7 | {audit.get('status', '\u672a\u6267\u884c')}\uff1b{audit.get('automated_checks', {}).get('cells_failed', 0)} \u683c\u672a\u901a\u8fc7\uff1brun_id={ab.get('run_id') or '\u65e0'} |",
        f"| \u9875\u7801\u6838\u9a8c | {page['manually_verified_pages']} / {page['total_pdf_pages']} \u9875 | {page['pending_manual_pages']} \u9875\u5f85\u4eba\u5de5\u6838\u9a8c |", "",
        "## \u56fe\u8868\u4e00\u89c8", "",
        "| \u56fe\u53f7 | \u5185\u5bb9 | N | \u72b6\u6001 |", "|---|---|---:|---|",
        f"| F01 | \u9996\u6279\u9898\u5e93 | {bank['question_count']} | \u6750\u6599\u6784\u6210 |",
        f"| F02 | A/B \u771f\u5b9e\u9a8c\u6536 | {ab['planned_cells']} \u8ba1\u5212\u683c | {ab['completed_cells']} \u5df2\u5b8c\u6210 |",
        f"| F03 | \u6559\u6750\u9875\u7801 | {page['total_pdf_pages']} | {page['manually_verified_pages']} \u9875\u5df2\u6838\u9a8c |",
        f"| F04 | \u5f00\u53d1\u6784\u9020\u6848\u4f8b | {cases['total']} | \u975e\u771f\u4eba\u53c2\u4e0e\u8005 |",
        f"| F05 | A/B \u8fd0\u884c\u6982\u89c8 | {ab['completed_cells']} \u683c / {perf['latency_ms']['n']} \u56de\u5408 | {audit.get('status', '\u672a\u6267\u884c')} |", "",
    ]
    desc = descriptions(snapshot)
    for figure_id, stem in FIGURES.items():
        numerator, denominator = {
            "F01": (bank["question_count"], bank["question_count"]),
            "F02": (ab["completed_cells"], ab["planned_cells"]),
            "F03": (page["total_pdf_pages"], page["total_pdf_pages"]),
            "F04": (cases["total"], cases["total"]),
            "F05": (ab["completed_cells"], ab["planned_cells"]),
        }[figure_id]
        lines += [
            f"## {figure_id} {stem}", "",
            f"![{figure_id}](figures/{stem}.png)", "",
            desc[figure_id][0], "",
            f"- {desc[figure_id][1]}",
            f"- \u89e3\u91ca\uff1a{desc[figure_id][2]}",
            f"- N={numerator} / \u603b\u5206\u6bcd={denominator}\uff1brun_id={ab['run_id'] or '\u65e0'}\u3002",
            f"- \u6570\u636e\uff1a[figure-data.csv](source-data/figure-data.csv)\uff0c[experiment-snapshot.json](source-data/experiment-snapshot.json)\u3002\u4ee3\u7801 [generate_charts.R](generate_charts.R)\uff1b\u73af\u5883 {runtime}\u3002",
            f"- \u8f93\u51fa\uff1a[PNG](figures/{stem}.png) \u00b7 [SVG](figures/{stem}.svg) \u00b7 [PDF](figures/{stem}.pdf) \u00b7 [TIFF](figures/{stem}.tiff)\u3002", "",
        ]
        if figure_id == "F01":
            lines += ["| \u6a21\u5757 | \u9898\u6570 | \u786e\u5b9a\u6027 | \u4eba\u5de5/\u672a\u542f\u7528 | \u5f15\u7528\u5f85\u6838 |", "|---|---:|---:|---:|---:|"]
            for module, counts in sorted(bank["review_by_module"].items()):
                lines.append(f"| {module} | {sum(counts.values())} | {counts.get('deterministic', 0)} | {counts.get('manual', 0)} | {counts.get('reference_pending', 0)} |")
            lines.append("")
        if figure_id == "F04":
            lines += ["| \u6848\u4f8b\u7c7b\u522b | \u6570\u91cf |", "|---|---:|"]
            for name, count in sorted(cases["categories"].items()):
                lines.append(f"| {CASE_NAMES.get(name, name)} | {count} |")
            lines.append("")
        if figure_id == "F05":
            lines += [
                "- \u9762\u677f a\uff1aA/B \u6848\u4f8b\u5b8c\u6210\u683c\uff1b\u9762\u677f b\uff1a\u6bcf\u56de\u5408\u7aef\u5230\u7aef\u5ef6\u8fdf\uff1b\u9762\u677f c\uff1a\u6bcf\u56de\u5408 Token \u603b\u91cf\u3002",
                f"- \u6027\u80fd\u6e90\u6570\u636e\uff1a[turn-performance.csv](source-data/turn-performance.csv)\uff1b\u4e0d\u542b\u7528\u6237\u63d0\u95ee\u6216\u6a21\u578b\u56de\u7b54\u6b63\u6587\u3002\u8bc4\u4f30\u660e\u7ec6\uff1a[acceptance-audit.json](source-data/acceptance-audit.json)\u3002",
                f"- \u5206\u7ec4\u5ef6\u8fdf\u548c Token \u7edf\u8ba1\uff1a{snapshot.get('performance_by_group', {})}\u3002",
                "- \u62fc\u56fe\u4f7f\u7528 ggplot2 + patchwork\uff1b\u53ea\u8868\u793a\u5de5\u7a0b\u8fd0\u884c\u6307\u6807\uff0c\u4e0d\u8868\u793a\u5b66\u4e60\u6210\u6548\u3002", "",
            ]
    token = perf["token_usage"]
    latency = perf["latency_ms"]
    findings_by_cell = {
        (str(item.get("case_id") or ""), str(item.get("group") or "")): item
        for item in audit.get("findings", []) if item.get("case_id")
    }
    failure_rows = [
        f"| {cell.get('case_id')} | {cell.get('group')} | {cell.get('failure_reason_code') or '\u2014'} | "
        f"{', '.join(findings_by_cell.get((str(cell.get('case_id') or ''), str(cell.get('group') or '')), {}).get('failed_checks', [])) or 'terminal_or_trace'} |"
        for cell in audit.get("cells", [])
        if cell.get("status") == "failed" or cell.get("checks", {}).get("cell_completed") is False
    ] or ["| \u2014 | \u2014 | \u2014 | \u65e0 |"]
    lines += [
        "## \u6027\u80fd\u6307\u6807", "",
        "| \u6307\u6807 | N | \u6570\u503c | \u5355\u4f4d |", "|---|---:|---:|---|",
        f"| \u5ef6\u8fdf\u4e2d\u4f4d\u6570 | {latency['n']} | {fmt(latency['median'])} | ms / \u811a\u672c\u5316\u56de\u5408 |",
        f"| \u5ef6\u8fdf P95 | {latency['n']} | {fmt(latency['p95'])} | ms / \u811a\u672c\u5316\u56de\u5408 |",
        f"| Token \u7528\u91cf | {token['n']} | {fmt(token['total_tokens'])} | prompt + completion |",
        f"| \u6a21\u578b\u8c03\u7528 | {fmt(perf['model_calls'])} | \u6b21 |",
        f"| \u4f30\u7b97\u8d39\u7528 | {token['n']} | {fmt(perf.get('estimated_cost'))} | {perf.get('cost_reason', '\u672a\u914d\u7f6e')} |", "",
        "## M1 \u81ea\u52a8\u9a8c\u6536\u672a\u901a\u8fc7\u683c", "",
        "\u5b8c\u6574 trace_id \u4fdd\u5b58\u5728\u672c\u673a\u8fd0\u884c\u4e8b\u4ef6\u5e93\u4e0e\u672c\u5730\u9a8c\u6536\u6e05\u5355\uff1b\u4e0b\u8868\u53ea\u5305\u542b\u8131\u654f\u539f\u56e0\u7801\u548c\u68c0\u67e5\u540d\u3002", "",
        "| \u6848\u4f8b | \u7ec4 | \u539f\u56e0\u7801 | \u5931\u8d25\u68c0\u67e5 |", "|---|---|---|---|",
        *failure_rows, "",
        acceptance_qa_status_line(snapshot), "",
        "## \u72ec\u7acb\u6545\u969c\u4e0e\u56fe\u8868 QA", "",
        f"- {automation_validation_line(snapshot)}",
        f"- \u6545\u969c\u6ce8\u5165 {validation.get('fault_injection', {}).get('status', 'not_run')}\uff08{validation.get('fault_injection', {}).get('tests_passed', 0)} \u9879\u6d4b\u8bd5\uff09\uff1bProvider \u9519\u8bef\u3001\u8d85\u65f6\u3001\u53d6\u6d88\u3001Gateway \u6062\u590d\u548c\u8131\u654f\u5bfc\u51fa\u5747\u7528\u72ec\u7acb\u6545\u969c\u6ce8\u5165\u9a8c\u8bc1\uff0c\u4e0d\u6df7\u5165\u771f\u5b9e Provider \u7ed3\u679c\u3002",
        f"- patchwork \u9762\u677f\u5bf9\u9f50 {validation.get('figures', {}).get('panel_alignment', {}).get('status', 'not_run')}\uff1bPDF \u6587\u5b57\u626b\u63cf {validation.get('figures', {}).get('pdf_text_audit', {}).get('status', 'not_run')}\uff1bPDF \u78b0\u649e\u626b\u63cf {validation.get('figures', {}).get('pdf_collision_audit', {}).get('status', 'not_run')}\u3002", "",
        "## \u8bfe\u7a0b\u8986\u76d6\u548c\u6570\u636e\u6765\u6e90", "",
        f"- \u8bfe\u7a0b {snapshot['course_coverage'].get('course_id')} \u5171 {snapshot['course_coverage'].get('concept_count')} \u4e2a\u77e5\u8bc6\u70b9\uff1b\u72b6\u6001\u5206\u5e03 {dict(Counter(row.get('coverage_status', 'unknown') for row in snapshot['course_coverage'].get('coverage', [])))}\u3002",
        f"- \u672a\u8986\u76d6\u77e5\u8bc6\u70b9\uff1a{', '.join(uncovered_concepts) if uncovered_concepts else '\u65e0'}\u3002",
        f"- \u9898\u96c6 SHA-256 {snapshot['source_hashes'].get('question_collection_pdf_sha256')}\u3002",
        f"- \u6559\u6750 SHA-256 {snapshot['source_hashes'].get('companion_textbook_pdf_sha256')}\u3002",
        "- \u9898\u96c6\u6b63\u6587\u3001\u7b54\u6848\u3001OCR \u548c\u539f\u59cb\u6a21\u578b\u8f93\u51fa\u4fdd\u5b58\u5728\u672c\u673a\u5f00\u53d1\u76ee\u5f55\uff1b\u4ed3\u5e93\u4ec5\u4fdd\u5b58\u805a\u5408\u6570\u636e\u548c\u5143\u6570\u636e\u3002\u5185\u90e8\u6d4b\u8bd5\u8bf4\u660e\u4e0d\u7b49\u540c\u4e8e\u6559\u6750\u8bb8\u53ef\u6838\u9a8c\u3002",
        "- \u6e90\u8bf4\u660e\u548c FAIR \u5143\u6570\u636e\u89c1 [\u6570\u636e\u4e0e\u6765\u6e90\u8bf4\u660e](\u6570\u636e\u4e0e\u6765\u6e90\u8bf4\u660e.md)\u3002", "",
        "## \u4e0b\u4e00\u6b65", "",
        *[f"- {item}" for item in snapshot.get("next_steps", [])], "",
        "## \u7248\u672c\u548c\u590d\u73b0", "",
        f"- commit {code.get('commit')}\uff1bworking tree dirty={code.get('working_tree_dirty')}\uff1b\u72b6\u6001\u6458\u8981 SHA-256 {code.get('working_tree_status_sha256')}\u3002",
        f"- \u5de5\u4f5c\u6811\u5dee\u5f02\u6307\u7eb9 {code.get('working_tree_diff_sha256')}\u3002",
        f"- \u9898\u5e93 {bank['version']}\uff1b\u6848\u4f8b {cases['version']}\uff1b\u7b56\u7565 {snapshot.get('runtime_configuration', {}).get('policy_version')}\uff1b\u6a21\u578b {snapshot.get('runtime_configuration', {}).get('model') or '\u672a\u914d\u7f6e'}\u3002",
        f"- \u9a8c\u8bc1\u72b6\u6001 {validation.get('status', '\u672a\u8bb0\u5f55')}\uff1bDOCX \u89c6\u89c9\u68c0\u67e5 {validation.get('docx_visual_qa', '\u672a\u8bb0\u5f55')}\uff1b\u8be6\u60c5 [validation.json](source-data/validation.json)\u3002",
        "- \u4f7f\u7528 deepprof login \u624b\u52a8\u914d\u7f6e Provider\uff1b\u9996\u6b21\u914d\u7f6e\u53ef\u89c1\u8f93\u5165 API Key\uff0c\u540e\u7eed\u9ed8\u8ba4\u9690\u85cf\uff1bdeepprof doctor \u786e\u8ba4\u63a2\u6d4b\u540e\u6267\u884c deepprof acceptance --live\uff0c\u5b8c\u6210\u540e\u6267\u884c deepprof report\u3002", "",
    ]
    return "\n".join(lines)


def set_font(run: Any, size: float = 10, bold: bool = False, color: str = "000000") -> None:
    run.font.name = "Microsoft YaHei"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_text(doc: Document, value: str, size: float = 10, bold: bool = False, color: str = "273746") -> Any:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.08
    set_font(p.add_run(value), size, bold, color)
    return p


def add_heading(doc: Document, value: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(5)
    set_font(p.add_run(value), 15 if level == 1 else 11, True, "000000")


def add_table(doc: Document, headers: list[str], values: list[list[str]], widths: list[float] | None = None) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade = OxmlElement("w:shd")
        shade.set(qn("w:fill"), "17324D")
        cell._tc.get_or_add_tcPr().append(shade)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_font(cell.paragraphs[0].add_run(value), 9, True, "FFFFFF")
    for row_index, row in enumerate(values):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "F1F5F8")
                cell._tc.get_or_add_tcPr().append(shade)
            cell.paragraphs[0].paragraph_format.space_after = Pt(1)
            set_font(cell.paragraphs[0].add_run(str(value)), 9, False)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_figure(doc: Document, figure_id: str) -> None:
    stem = FIGURES[figure_id]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(1)
    p.add_run().add_picture(str(DOCS / "figures" / f"{stem}.png"), width=Inches(6.2))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(3)
    set_font(caption.add_run(f"{figure_id}  {stem}"), 9, True, "5D6870")


def build_docx(snapshot: dict[str, Any]) -> None:
    bank = snapshot["question_bank"]
    ab = snapshot["real_ab_acceptance"]
    perf = snapshot["performance"]
    pages = snapshot["textbook_page_audit"]
    cases = snapshot["constructed_cases"]
    validation = snapshot.get("validation", {})
    coverage = Counter(row.get("coverage_status", "unknown") for row in snapshot["course_coverage"].get("coverage", []))
    uncovered_concepts = [row.get("concept_id", "") for row in snapshot["course_coverage"].get("coverage", [])
                          if row.get("coverage_status") in {"uncovered", "not_covered"}]
    audit = snapshot.get("acceptance_audit") or {}
    runtime = plot_versions()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin, section.bottom_margin = Inches(.55), Inches(.55)
    section.left_margin, section.right_margin = Inches(.65), Inches(.65)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("DeepProf  M0 / M1  "), 8, False, "718096")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    title = doc.add_paragraph(style="Title")
    title.paragraph_format.space_before = Pt(20)
    title.paragraph_format.space_after = Pt(4)
    set_font(title.add_run("DeepProf  \u9a8c\u6536\u4e0e\u5b9e\u9a8c\u56fe\u8868\u5e93"), 25, True, "000000")
    add_text(doc, "M0 \u6750\u6599\u4e0e M1 \u6280\u672f\u9a8c\u6536\u57fa\u7ebf", 12, True, "000000")
    add_text(doc, f"\u5feb\u7167 {snapshot['generated_at']}  \u9898\u5e93 {bank['version']}  \u6848\u4f8b {cases['version']}", 9, False, "5D6870")
    add_heading(doc, "\u9a8c\u6536\u6458\u8981")
    add_table(doc, ["\u9879\u76ee", "\u8bb0\u5f55", "\u5224\u8bfb"], [
        ["M0 \u5de5\u7a0b\u6750\u6599", f"{bank['question_count']} \u9898 / {cases['total']} \u4f8b", f"{bank['verified_source_count']} \u9898\u6765\u6e90\u5df2\u590d\u6838\uff1b{bank['pending_reference_count']} \u9898\u5f15\u7528\u5f85\u6838"],
        ["\u786e\u5b9a\u6027\u8bc4\u5206", f"{bank['deterministic_count']} / {bank['question_count']}", "\u4ec5\u542f\u7528\u6765\u6e90\u548c\u89c4\u5219\u90fd\u901a\u8fc7\u6838\u9a8c\u7684\u9898\u76ee"],
        ["\u6559\u5e08\u5ba1\u6838", str(bank["teacher_approval"]), "\u4fdd\u7559\u5f85\u5ba1"],
        ["M1 \u771f\u5b9e A/B", f"{ab['completed_cells']} / {ab['planned_cells']} \u683c", f"{ab.get('provider_profile') or '\u672a\u914d\u7f6e'} / {ab.get('model') or '\u672a\u914d\u7f6e'}\uff1b\u81ea\u52a8\u6848\u4f8b\u6838\u9a8c {audit.get('status', '\u672a\u6267\u884c')}"],
        ["M1 \u81ea\u52a8\u884c\u4e3a", f"{audit.get('automated_checks', {}).get('cells_passed', 0)} / {ab['planned_cells']} \u683c", f"{audit.get('status', '\u672a\u6267\u884c')}\uff1b{audit.get('automated_checks', {}).get('cells_failed', 0)} \u683c\u672a\u901a\u8fc7"],
        ["\u6559\u6750\u9875\u7801", f"{pages['manually_verified_pages']} / {pages['total_pdf_pages']} \u9875", f"{pages['pending_manual_pages']} \u9875\u5f85\u6838"],
    ], [1.15, 1.55, 3.5])
    add_heading(doc, "\u6570\u636e\u8fb9\u754c", 2)
    add_text(doc, f"\u672c\u6b21 A/B \u4f7f\u7528\u771f\u5b9e Provider {ab.get('provider_profile') or '\u672a\u914d\u7f6e'} / {ab.get('model') or '\u672a\u914d\u7f6e'}\uff0c\u8bb0\u5f55 {ab['completed_cells']}/{ab['planned_cells']} \u683c\u4e0e {perf['model_calls']} \u6b21\u6a21\u578b\u8c03\u7528\u3002\u8fd0\u884c\u6307\u6807\u662f\u6784\u9020\u6848\u4f8b\u7684\u5de5\u7a0b\u6d4b\u91cf\uff0c\u4e0d\u662f\u5b66\u4e60\u6548\u679c\u3002\u6559\u5e08\u8bc4\u5206\u3001\u8bed\u4e49\u6cc4\u9732\u548c\u5f15\u7528\u652f\u6301\u5ea6\u4ecd\u5f85\u4eba\u5de5\u8bc4\u4f30\u3002")
    doc.add_page_break()

    add_heading(doc, "F01  \u9996\u6279\u9898\u5e93\u6784\u6210")
    add_figure(doc, "F01")
    add_text(doc, "\u5206\u5b50/\u5206\u6bcd\uff1a\u5404\u6a21\u5757\u5404\u72b6\u6001\u9898\u6570 / \u8be5\u6a21\u5757\u9898\u6570\uff1b\u5355\u4f4d\u9053\uff1bN=30\u3002", 10)
    module_rows = [[module, str(sum(counts.values())), str(counts.get("deterministic", 0)), str(counts.get("manual", 0)), str(counts.get("reference_pending", 0))] for module, counts in sorted(bank["review_by_module"].items())]
    add_table(doc, ["\u6a21\u5757", "\u9898\u6570", "\u786e\u5b9a\u6027", "\u4eba\u5de5/\u5f85\u542f\u7528", "\u5f15\u7528\u5f85\u6838"], module_rows, [1.2, .65, 1.0, 1.65, 1.0])
    add_text(doc, "DS-GRAPH-04-Q7.4 \u7684\u9898\u9762\u3001\u7b54\u6848\u548c\u4e66\u5185\u56fe\u9875\u5f15\u7528\u5df2\u590d\u6838\uff1b\u8fd9\u9053\u5f00\u653e\u9898\u4fdd\u6301\u4eba\u5de5\u8bc4\u5206\uff0c\u81ea\u52a8\u8bc4\u5206\u7981\u7528\uff0c\u9898\u5e93\u6559\u5e08\u5ba1\u6838\u4ecd\u5f85\u5b8c\u6210\u3002\u626b\u63cf\u9898\u9762\u548c\u7b54\u6848\u4ec5\u4fdd\u5b58\u5728\u672c\u673a\u6d4b\u8bd5\u76ee\u5f55\u3002")
    doc.add_page_break()

    add_heading(doc, "F02  A\u7ec4\u4e0eB\u7ec4\u9a8c\u6536\u8fdb\u5ea6\u53ca\u6027\u80fd\u6570\u636e")
    add_figure(doc, "F02")
    add_text(doc, f"\u8bbe\u8ba1\u4e3a 40 \u4e2a\u51bb\u7ed3\u6848\u4f8b\u5728 A\u3001B \u4e24\u7ec4\u5404\u8fd0\u884c\u4e00\u6b21\uff0c\u517180 \u683c\u3002\u5f53\u524d\u8bb0\u5f55 {ab['completed_cells']} \u683c\uff0cProvider {ab.get('provider_profile') or '\u672a\u914d\u7f6e'} / {ab.get('model') or '\u672a\u914d\u7f6e'}\uff0cFake Provider={ab.get('fake_provider_used', False)}\u3002")
    add_text(doc, f"\u5206\u5b50/\u5206\u6bcd\uff1a{ab['completed_cells']}/{ab['planned_cells']} \u683c\uff1bN={ab['planned_cells']}\uff1brun_id={ab['run_id'] or '\u65e0'}\uff1b\u81ea\u52a8\u884c\u4e3a\u6838\u9a8c={audit.get('status', '\u672a\u6267\u884c')}\u3002", 10)
    add_table(doc, ["\u6307\u6807", "N", "\u6570\u503c", "\u5355\u4f4d"], [
        ["\u56de\u5408\u5ef6\u8fdf\u4e2d\u4f4d\u6570", str(perf["latency_ms"]["n"]), fmt(perf["latency_ms"]["median"]), "ms / \u811a\u672c\u5316\u6a21\u578b\u56de\u5408"],
        ["\u56de\u5408\u5ef6\u8fdf P95", str(perf["latency_ms"]["n"]), fmt(perf["latency_ms"]["p95"]), "ms / \u811a\u672c\u5316\u6a21\u578b\u56de\u5408"],
        ["Token \u7528\u91cf", str(perf["token_usage"]["n"]), fmt(perf["token_usage"]["total_tokens"]), "prompt + completion"],
        ["\u6a21\u578b\u8c03\u7528", fmt(perf["model_calls"]), fmt(perf["model_calls"]), "\u6b21"],
        ["\u4f30\u7b97\u8d39\u7528", str(perf["token_usage"]["n"]), fmt(perf.get("estimated_cost")), str(perf.get("cost_reason") or "")],
    ], [1.5, .6, .8, 3.0])
    add_text(doc, f"\u6559\u5e08\u8bc4\u5206\u3001\u8bed\u4e49\u6cc4\u9732\u3001\u5f15\u7528\u652f\u6301\u5ea6\u5f85\u4eba\u5de5\u8bc4\u4f30\uff1b\u6545\u969c\u6ce8\u5165={snapshot['manual_evaluation'].get('fault_injection', 'not_run')}\uff1b\u8fdb\u7a0b\u6062\u590d={snapshot['manual_evaluation'].get('process_recovery', 'not_run')}\u3002")
    doc.add_page_break()

    add_heading(doc, "F03  \u914d\u5957\u6559\u6750\u9875\u7801\u6838\u9a8c")
    add_figure(doc, "F03")
    page_audit_rows = [
        ["PDF \u603b\u9875\u6570", str(pages["total_pdf_pages"]), "", ""],
        ["\u4eba\u5de5\u5df2\u6838\u9a8c", str(pages["manually_verified_pages"]), "", ""],
        ["\u5f85\u6838", str(pages["pending_manual_pages"]), "", ""],
    ]
    page_audit_rows.extend([
        ["\u4eba\u5de5\u6620\u5c04", str(item.get("pdf_page") or ""),
         str(item.get("printed_page") or "\u672a\u8bc6\u522b"), str(item.get("chapter_heading") or "")]
        for item in pages.get("verified_mappings", [])
    ])
    page_audit_rows.append(["\u7edf\u4e00\u504f\u79fb", "", "", "\u672a\u5047\u8bbe"])
    add_table(doc, ["\u6838\u9a8c\u9879", "PDF \u9875", "\u4e66\u5185\u9875", "\u7ae0\u8282 / \u6807\u8bc6"],
              page_audit_rows, [1.35, .75, .85, 3.2])
    add_text(doc, "\u5f15\u7528\u9875\u7801\u9700\u9010\u9875\u4eba\u5de5\u6838\u5bf9\uff1b\u672a\u786e\u8ba4\u9875\u4fdd\u7559\u5f85\u6838\u6807\u8bb0\u3002")
    doc.add_page_break()

    add_heading(doc, "F04  \u51bb\u7ed3\u7684\u5f00\u53d1\u6784\u9020\u6848\u4f8b")
    add_figure(doc, "F04")
    case_rows = [[CASE_NAMES.get(name, name), str(count)] for name, count in sorted(cases["categories"].items())]
    add_table(doc, ["\u6848\u4f8b\u7c7b\u522b", "\u6570\u91cf"], case_rows, [4.4, 1.75])
    add_text(doc, f"\u7248\u672c {cases['version']}\uff0c\u6bcf\u7c7b 5 \u4f8b\uff0c\u6848\u4f8b\u6e05\u5355\u4fdd\u5b58\u5728 evaluation/dev_cases.json\u3002\u8fd9\u4e9b\u8f93\u5165\u4e0d\u662f\u5b66\u751f\u6570\u636e\u3002")
    doc.add_page_break()

    add_heading(doc, "F05  A/B \u771f\u5b9e\u8fd0\u884c\u6982\u89c8")
    add_figure(doc, "F05")
    add_text(doc, f"N={ab['completed_cells']}/{ab['planned_cells']} \u683c\uff1brun_id={ab['run_id'] or '\u65e0'}\uff1b\u5ef6\u8fdf N={perf['latency_ms']['n']}\uff0cToken N={perf['token_usage']['n']}\uff0c\u6a21\u578b\u8c03\u7528={perf['model_calls']} \u6b21\u3002Token \u5305\u542b\u5931\u8d25\u56de\u5408\u7684\u53ef\u7528\u7528\u91cf\uff1b\u4e0d\u542b\u63d0\u95ee\u6216\u6a21\u578b\u6b63\u6587\u3002", 9)
    group_rows = []
    for group in ("A", "B"):
        group_perf = snapshot.get("performance_by_group", {}).get(group, {})
        latency_group = group_perf.get("latency_ms", {})
        usage_group = group_perf.get("token_usage", {})
        counts = ab.get("groups", {}).get(group, {})
        group_rows.append([
            group, str(counts.get("completed", 0)), str(latency_group.get("n", 0)),
            fmt(latency_group.get("median")), fmt(latency_group.get("p95")),
            str(usage_group.get("n", 0)), fmt(usage_group.get("total_tokens")), str(group_perf.get("model_calls", 0)),
        ])
    add_table(doc, ["\u7ec4", "\u5b8c\u6210/40", "\u5ef6\u8fdfN", "\u4e2d\u4f4d ms", "P95 ms", "Token N", "Token", "\u8c03\u7528\u6b21\u6570"],
              group_rows, [.4, .7, .55, .85, .65, .6, 1.0, 1.0])
    add_text(doc, f"\u81ea\u52a8\u884c\u4e3a\u6838\u9a8c {audit.get('status', '\u672a\u6267\u884c')}\uff1b\u6240\u6709\u6559\u5b66\u54cd\u5e94\u4ecd\u9700\u6559\u5e08\u8bc4\u5206\u3002\u6e90\u6570\u636e\u89c1 source-data/turn-performance.csv\u3002")
    failed_case_ids = sorted({str(item.get("case_id")) for item in audit.get("findings", []) if item.get("case_id")})
    failed_runtime_cells = [cell for cell in audit.get("cells", []) if cell.get("failure_reason_code")]
    runtime_failure_text = ", ".join(
        f"{cell['case_id']} {cell['group']} ({cell['failure_reason_code']})" for cell in failed_runtime_cells
    ) or "0"
    add_table(doc, ["\u81ea\u52a8\u68c0\u67e5", "\u7ed3\u679c", "\u5355\u4f4d / \u5931\u8d25 case_id"], [
        ["\u7ec8\u6001\u683c\u6570", str(audit.get("matrix", {}).get("completed", 0)), "\u8ba1\u5212 80 \u683c"],
        ["\u89c4\u5219\u901a\u8fc7", str(audit.get("automated_checks", {}).get("cells_passed", 0)), "\u683c"],
        ["\u89c4\u5219\u672a\u901a\u8fc7", str(audit.get("automated_checks", {}).get("cells_failed", 0)), "\u683c"],
        ["\u8bc1\u636e\u5b9a\u4f4d/\u7f3a\u53e3\u5931\u8d25", str(audit.get("automated_checks", {}).get("evidence_locator_or_gap_failures", 0)), "\u683c"],
        ["\u9884\u671f\u52a8\u4f5c\u5931\u8d25", str(audit.get("automated_checks", {}).get("expected_action_mismatches", 0)), ", ".join(failed_case_ids) or "0"],
        ["\u7ec8\u6001\u8fd0\u884c\u5931\u8d25", str(len(failed_runtime_cells)), runtime_failure_text],
    ], [2.2, 1.0, 2.9])
    add_text(doc, "\u89c4\u5219\u6838\u67e5\u4e0d\u4ee3\u66ff\u6559\u5e08\u8bc4\u5206\uff1b\u6559\u5e08\u8bc4\u5206\u3001\u8bed\u4e49\u6cc4\u9732\u4e0e\u5f15\u7528\u652f\u6301\u5ea6\u72b6\u6001\u5355\u72ec\u4fdd\u7559\u4e3a\u5f85\u8bc4\u3002", 9)
    doc.add_page_break()

    add_heading(doc, "\u6570\u636e\u6765\u6e90\u8986\u76d6\u4e0e\u590d\u73b0")
    add_heading(doc, "\u8bfe\u7a0b\u8986\u76d6", 2)
    add_text(doc, f"\u5171 {snapshot['course_coverage'].get('concept_count', 0)} \u4e2a\u8bfe\u7a0b\u77e5\u8bc6\u70b9\uff0c\u6620\u5c04\u72b6\u6001 {dict(coverage)}\u3002\u672a\u8986\u76d6\u9879\u4fdd\u7559\u5728\u8868\u4e2d\uff0c\u4e0d\u7528\u65e0\u5173\u9898\u76ee\u586b\u5145\u3002")
    add_text(doc, f"\u672a\u8986\u76d6\u77e5\u8bc6\u70b9\uff1a{', '.join(uncovered_concepts) if uncovered_concepts else '\u65e0'}\u3002", 9)
    add_heading(doc, "\u6765\u6e90\u548c\u8303\u56f4", 2)
    add_text(doc, "\u9898\u96c6\u3001\u7b54\u6848\u548c\u6559\u6750\u9875\u9762\u4ec5\u5b58\u5728\u672c\u673a\u5f00\u53d1\u6d4b\u8bd5\u76ee\u5f55\uff1b\u4ed3\u5e93\u4ec5\u4fdd\u5b58\u805a\u5408\u6570\u636e\u3001\u5143\u6570\u636e\u548c\u811a\u672c\u3002\u7528\u6237\u9650\u5b9a\u5185\u90e8\u6d4b\u8bd5\uff0c\u4e0d\u8868\u793a\u72ec\u7acb\u6559\u6750\u8bb8\u53ef\u6838\u9a8c\u3002")
    add_table(doc, ["\u6765\u6e90", "SHA-256"], [
        ["\u9898\u96c6 PDF", str(snapshot["source_hashes"].get("question_collection_pdf_sha256"))],
        ["\u914d\u5957\u6559\u6750 PDF", str(snapshot["source_hashes"].get("companion_textbook_pdf_sha256"))],
        ["\u6848\u4f8b\u5143\u6570\u636e", str(snapshot["source_hashes"].get("constructed_case_file_sha256"))],
    ], [2.0, 4.15])
    add_heading(doc, "\u81ea\u52a8\u5316\u68c0\u67e5", 2)
    add_text(doc, f"Python {validation.get('python', {}).get('passed', '\u672a\u8bb0\u5f55')} \u9879\u901a\u8fc7 / {validation.get('python', {}).get('skipped', '\u2014')} \u9879\u8df3\u8fc7\uff1bCLI {validation.get('cli', {}).get('tests_passed', '\u672a\u8bb0\u5f55')} \u9879\u901a\u8fc7\uff1b\u7c7b\u578b\u68c0\u67e5 {validation.get('cli', {}).get('typecheck', '\u672a\u8bb0\u5f55')}\u3002")
    fault = validation.get("fault_injection", {})
    figure_qa = validation.get("figures", {})
    add_text(doc, f"\u72ec\u7acb\u6545\u969c\u6ce8\u5165 {fault.get('status', 'not_run')}\uff08{fault.get('tests_passed', 0)} \u9879\uff09\uff1b\u8fdb\u7a0b\u6062\u590d {validation.get('process_recovery', 'not_run')}\u3002\u56fe\u8868\u9762\u677f\u5bf9\u9f50 {figure_qa.get('panel_alignment', {}).get('status', 'not_run')}\uff1bPDF \u6587\u5b57\u626b\u63cf {figure_qa.get('pdf_text_audit', {}).get('status', 'not_run')}\uff1bPDF \u78b0\u649e\u626b\u63cf {figure_qa.get('pdf_collision_audit', {}).get('status', 'not_run')}\u3002", 9)
    add_text(doc, f"DOCX \u89c6\u89c9\u68c0\u67e5 {validation.get('docx_visual_qa', '\u672a\u8bb0\u5f55')}\uff1b{validation.get('docx_renderer', '')}", 9)
    add_text(doc, f"commit {snapshot['code'].get('commit')}\uff1b\u5de5\u4f5c\u6811\u5dee\u5f02\u6307\u7eb9 {snapshot['code'].get('working_tree_diff_sha256')}\u3002\u4f5c\u56fe\u73af\u5883 {runtime}\u3002", 9)
    add_heading(doc, "\u4e0b\u4e00\u6b65", 2)
    for item in snapshot.get("next_steps", []):
        add_text(doc, f"\u2022 {item}", 9)

    doc.core_properties.title = "DeepProf \u9a8c\u6536\u4e0e\u5b9e\u9a8c\u56fe\u8868\u5e93"
    doc.core_properties.subject = "M0 engineering and M1 technical acceptance"
    doc.core_properties.author = "DeepProf"
    doc.core_properties.comments = f"{snapshot['snapshot_version']} {datetime.now(timezone.utc).isoformat()}"
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_PATH)


def main() -> None:
    global DOCX_PATH
    parser = argparse.ArgumentParser()
    parser.add_argument("--markdown-only", action="store_true")
    parser.add_argument("--docx-only", action="store_true")
    args = parser.parse_args()
    if args.markdown_only and args.docx_only:
        raise SystemExit("Select at most one output restriction.")
    snapshot = load_snapshot()
    rows = read_csv()
    verify(snapshot, rows)
    if not args.markdown_only:
        try:
            build_docx(snapshot)
        except PermissionError:
            # Office applications may hold the canonical report open. Preserve
            # that live document and write a snapshot-named copy instead.
            stamp = snapshot["generated_at"].replace(":", "").replace("-", "").replace("T", "-").split(".")[0]
            DOCX_PATH = DOCS / f"图表库-快照-{stamp}.docx"
            build_docx(snapshot)
            print(f"Canonical DOCX is open; snapshot copy written: {DOCX_PATH}")
        else:
            print(f"DOCX written: {DOCX_PATH}")
    if not args.docx_only:
        markdown = make_markdown(snapshot, rows).replace("](图表库.docx)", f"]({DOCX_PATH.name})")
        MD_PATH.write_text(markdown, encoding="utf-8")
        print(f"Markdown written: {MD_PATH}")


if __name__ == "__main__":
    main()
